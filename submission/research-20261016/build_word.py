"""Build editable Word from the canonical manuscript and IEEE Word template."""
from pathlib import Path
import re,subprocess,copy
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
P=Path(__file__).resolve().parent;tex=(P/'paper.tex').read_text()
keys=re.findall(r'\\bibitem\{([^}]+)\}',tex);nums={k:i+1 for i,k in enumerate(keys)}
body=tex[tex.index(r'\begin{abstract}'):tex.index(r'\balance')]
body=re.sub(r'\\cite\{([^}]+)\}',lambda m:', '.join('['+str(nums[k])+']' for k in m[1].split(',')),body)
for key,val in [('fig:architecture','1'),('fig:delta','2'),('tab:setup','I'),('tab:results','II')]:body=body.replace(r'\ref{'+key+'}',val)
body=re.sub(r'\\label\{[^}]+\}','',body)
body=body.replace(r'\begin{abstract}',r'\noindent\textbf{Abstract—} ').replace(r'\end{abstract}','\n')
body=body.replace(r'\begin{IEEEkeywords}',r'\noindent\textbf{Index Terms—} ').replace(r'\end{IEEEkeywords}','\n')
for name in ['architecture','rag-delta']:body=body.replace(r'\includegraphics[width=\columnwidth]{figures/'+name+'.pdf}',r'\includegraphics[width=3.4in]{figures/'+name+'.png}')
body=body.replace(r'\caption{Implemented investigation architecture.',r'\caption{Fig. 1. Implemented investigation architecture.').replace(r'\caption{MAP differences produced by LLM reranking.',r'\caption{Fig. 2. MAP differences produced by LLM reranking.')
bib=tex[tex.index(r'\begin{thebibliography}'):tex.index(r'\end{thebibliography}')]
entries=re.split(r'\\bibitem\{[^}]+\}',bib)[1:]
body+='\n'+r'\section*{References}'+'\n'+'\n\n'.join('['+str(i+1)+'] '+entry.strip()for i,entry in enumerate(entries))
work=Path('/tmp/vigilant-word-build');work.mkdir(exist_ok=True)
(work/'conversion.tex').write_text('\\documentclass{article}\n\\begin{document}\n'+body+'\n\\end{document}\n')
subprocess.run(['pandoc',str(work/'conversion.tex'),'-f','latex','-t','docx','--reference-doc',str(P/'IEEE-Word-Template.docx'),'--resource-path',str(P),'-o',str(work/'converted.docx')],check=True)
d=Document(work/'converted.docx');template=Document(P/'IEEE-Word-Template.docx')
for p in d.paragraphs:
 ref=p._p.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
 if ref is not None:
  old=ref.get(qn('w:val'))
  mapping={'FirstParagraph':'BodyText','TableCaption':'tablehead','ImageCaption':'figurecaption','CaptionedFigure':'BodyText'}
  if old in mapping:ref.set(qn('w:val'),mapping[old])
bodysec=copy.deepcopy(next(s._sectPr for s in template.sections if s._sectPr.find(qn('w:cols'))is not None and s._sectPr.find(qn('w:cols')).get(qn('w:num'))=='2'))
for el in list(bodysec):
 if el.tag in [qn('w:headerReference'),qn('w:footerReference'),qn('w:titlePg')]:bodysec.remove(el)
for el in list(d.element.body):
 if el.tag==qn('w:sectPr'):d.element.body.remove(el)
d.element.body.append(bodysec)
for name in ['Normal','Body Text','Heading 1','Heading 2','Abstract','Keywords','references','figure caption','table head','table copy']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.xpath('./w:pPr/w:numPr')):el.getparent().remove(el)
for name in ['Normal','Body Text']:
 st=d.styles[name];st.font.size=Pt(10);st.paragraph_format.space_after=Pt(0);st.paragraph_format.space_before=Pt(0);st.paragraph_format.line_spacing=.95;st.paragraph_format.first_line_indent=Inches(.12);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2']:
 st=d.styles[name];st.font.size=Pt(10);st.font.bold=False;st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True;st.paragraph_format.first_line_indent=Pt(0)
d.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.styles['Heading 2'].font.italic=True
d.styles['Heading 2'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
first=d.paragraphs[0]
title=first.insert_paragraph_before(re.search(r'\\title\{([^}]+)\}',tex)[1])
title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.paragraph_format.first_line_indent=Pt(0);title.paragraph_format.space_after=Pt(20);title.paragraph_format.line_spacing=1
for r in title.runs:r.font.name='Times New Roman';r.font.size=Pt(24)
author=first.insert_paragraph_before('Ankith Singh');author.alignment=WD_ALIGN_PARAGRAPH.CENTER;author.paragraph_format.first_line_indent=Pt(0);author.paragraph_format.space_after=Pt(28)
for r in author.runs:r.font.name='Times New Roman';r.font.size=Pt(11)
sect=copy.deepcopy(bodysec);sect.find(qn('w:cols')).set(qn('w:num'),'1');author._p.get_or_add_pPr().append(sect)
roman=['I','II','III','IV','V','VI','VII','VIII'];n=0;sub=0;references=False;setup_caption_seen=False
for p in d.paragraphs:
 t=p.text
 if p in [title,author]:continue
 if p.style.name=='Heading 1':
  if t=='References':p.text='References';references=True
  else:p.text=roman[n]+'. '+t;n+=1;sub=0
 elif p.style.name=='Heading 2':p.text=chr(65+sub)+'. '+t;sub+=1
 elif t.startswith('Abstract—'):
  p.style=d.styles['Abstract'];p.paragraph_format.line_spacing=.95;p.paragraph_format.space_after=Pt(5);p.paragraph_format.first_line_indent=Pt(0)
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=True
 elif t.startswith('Index Terms—'):
  p.paragraph_format.space_after=Pt(4);p.paragraph_format.first_line_indent=Pt(0)
  for r in p.runs:r.font.size=Pt(9);r.bold=True;r.italic=True
 elif references and t:
  p.style=d.styles['references'];p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(1);p.paragraph_format.first_line_indent=Inches(-.14);p.paragraph_format.left_indent=Inches(.14)
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
 elif t.startswith('Fig. '):
  p.style=d.styles['figure caption'];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1
  for r in p.runs:r.font.size=Pt(8)
 elif p.style.name=='table head':
  if 'Datasets' in t and setup_caption_seen:
   p._element.getparent().remove(p._element);continue
  if 'Datasets' in t:setup_caption_seen=True
  p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
  number='I'if 'Datasets' in t else 'II';p.text='TABLE '+number+'\n'+t
  for r in p.runs:r.font.size=Pt(8)
 if p._p.xpath('.//w:drawing'):
  p.paragraph_format.line_spacing=1.0
  p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
 if p._p.xpath('.//m:oMathPara'):
  p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6);p.paragraph_format.first_line_indent=Pt(0)
  p._p.get_or_add_pPr().append(OxmlElement('w:keepLines'))
  for r in p.runs:r.font.size=Pt(10)
for table in d.tables:
 pr=table._tbl.tblPr
 for el in list(pr):
  if el.tag==qn('w:tblStyle'):pr.remove(el)
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'5040')
 margins=OxmlElement('w:tblCellMar')
 for side in ['top','left','bottom','right']:
  m=OxmlElement('w:'+side);m.set(qn('w:w'),'35');m.set(qn('w:type'),'dxa');margins.append(m)
 pr.append(margins)
 table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 width=3.5;count=len(table.columns)
 widths=([1.5,.45,.5,.58,.47]if count==5 else [.85,2.65])
 for j,col in enumerate(table.columns):col.width=Inches(widths[j])
 for i,row in enumerate(table.rows):
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j])
   for p in cell.paragraphs:
    p.style=d.styles['table copy']
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Pt(0)
    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(8);run.bold=i==0
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shade)
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
for sec in d.sections:
 for h in [sec.header,sec.footer,sec.first_page_footer,sec.even_page_footer]:
  for p in h.paragraphs:p.text=''
d.core_properties.title=re.search(r'\\title\{([^}]+)\}',tex)[1];d.core_properties.author='Ankith Singh'
d.save(P/'VIGILANT-Research-Paper.docx')
print('Built editable Word manuscript.')
